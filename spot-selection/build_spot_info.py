"""Write the 16 Asakusa checkpoints in the Suisui Tabi spot-information format
(sheets 'spotsinfo' and 'spotsinfo - EN', same columns as suisuitabi_information_-_EN.xlsx),
plus a stamp-rally sheet and a notes sheet."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

rows = json.load(open('final.json'))['rows']
coords = {r['name']: (r['lat'], r['lng']) for r in rows if r['selected']}

HDR_JP = ['ID', '名前', '緯度', '経度', 'カテゴリ', 'ジャンル', 'タグ', 'ミッション', '地域', '営業時間', 'URL', 'おすすめポイント（コメント）']
HDR_EN = ['ID', 'Name', 'Latitude', 'Longtitude', 'Categories', 'Genre', 'Tag', 'Mission', 'Location', 'Business hours', 'URL', 'Recommendation points']
DAYS_JP = ['月曜日', '火曜日', '水曜日', '木曜日', '金曜日', '土曜日', '日曜日']
DAYS_EN = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
LOC_JP, LOC_EN = '上野・浅草・日暮里・谷根千', 'Ueno, Asakusa, Nippori, Yanesen'

TAG = {  # Japanese -> English, using the vocabulary already in the file
    '# 子供連れ': '# With children', '# ひとり旅': '# Solo travel', '# 女子旅': "# Girls' trip", '# デート': '# Date',
    '# 思い出': '# Memories', '# 映え': '# Stunning', '# 定番': '# Standard', '# 穴場': '# Hidden spots',
    '# 新発見': '# New discovery', '# 癒やし': '# Healing', '# 雨の日': '# Rainy day', '# 大人っぽい': '# Mature'}
GENRE = {'史跡・寺社': 'Historical sites/Temples and Shrines', '博物館・美術館': 'Museum/Art Museum', '公園': 'Park', '街並み': 'Townscape'}


def hours(open_, close):
    if open_ is None:
        return None, None
    if open_ == '24h':
        return (','.join(f'{d}: 24 時間営業' for d in DAYS_JP), ','.join(f'{d}: Open 24 hours' for d in DAYS_EN))
    h1, m1 = open_.split(':'); h2, m2 = close.split(':')
    jp = ','.join(f'{d}: {int(h1)}時{m1}分〜{int(h2)}時{m2}分' for d in DAYS_JP)
    en = ', '.join(f'{d}: {int(h1)}:{m1}-{int(h2)}:{m2}' for d in DAYS_EN)
    return jp, en


def gmaps(lat, lng):
    return f'https://www.google.com/maps/search/?api=1&query={lat:.7f},{lng:.7f}'


# (key in final.json, ID, JP name, EN name, category JP/EN, genre JP, tags JP, hours, URL or None, JP comment, EN comment)
SPOTS = [
    ('Sensō-ji', 1041, '浅草寺', 'Senso-ji Temple', ('観光', 'Tourism'), '史跡・寺社',
     ['# 子供連れ', '# ひとり旅', '# 女子旅', '# デート', '# 思い出', '# 映え', '# 定番'], ('6:30', '17:00'), 'https://www.senso-ji.jp/',
     '浅草のシンボルの雷門と書かれた赤い大提灯がある「浅草寺」では、観光客のフォトスポットとなっている人気の観光スポットです！また、昼は商店街として賑わい、夜はライトアップで綺麗なスカイツリーが見えるスポットです！\n\n是非一度、王道の観光スポット「浅草寺」に訪れてみるのはいかがでしょうか？',
     'Sensoji Temple is a popular tourist spot with a large red lantern that says Kaminarimon, the symbol of Asakusa, and is a popular photo spot for tourists! Also, it is a busy shopping street during the day, and at night you can see the beautiful Sky Tree lit up!\n\nWhy not visit the famous tourist spot "Sensoji Temple" at least once?'),
    ('Kaminarimon Gate', None, '雷門', 'Kaminarimon Gate', ('観光', 'Tourism'), '史跡・寺社',
     ['# 子供連れ', '# ひとり旅', '# デート', '# 思い出', '# 映え', '# 定番'], ('24h', None), None,
     '浅草寺の総門「雷門」は、大きな赤い提灯がかかる浅草のシンボルで、記念撮影の定番スポットです！門の左右には風神・雷神像が立ち、門をくぐると仲見世通りが浅草寺へと続きます。\n\n是非一度、浅草観光の出発点「雷門」に訪れてみるのはいかがでしょうか？',
     'Kaminarimon, the main gate of Senso-ji with its giant red lantern, is the symbol of Asakusa and a classic photo spot! Statues of the wind god and thunder god stand on either side, and Nakamise shopping street leads from the gate to the temple.\n\nWhy not start your Asakusa trip at "Kaminarimon"?'),
    ('Kappabashi Kitchen Supply Town', None, 'かっぱ橋道具街', 'Kappabashi Kitchen Town', ('観光', 'Tourism'), '街並み',
     ['# ひとり旅', '# 女子旅', '# デート', '# 思い出', '# 新発見', '# 定番'], None, 'https://www.kappabashi.or.jp/',
     '約800mの通りに調理器具や食器、包丁、食品サンプルなどの専門店が並ぶ「かっぱ橋道具街」は、プロの料理人から観光客まで楽しめるスポットです！食品サンプル作り体験ができるお店もあります。\n\n是非一度、日本の台所道具が揃う「かっぱ橋道具街」に訪れてみるのはいかがでしょうか？',
     'Kappabashi Kitchen Town is an 800 m street of shops selling cookware, tableware, knives and plastic food samples, loved by chefs and tourists alike! Some shops even offer food-sample-making experiences.\n\nWhy not visit "Kappabashi" to find Japanese kitchen tools?'),
    ('Imado Shrine', None, '今戸神社', 'Imado Shrine', ('観光', 'Tourism'), '史跡・寺社',
     ['# ひとり旅', '# 女子旅', '# デート', '# 映え', '# 穴場', '# 新発見'], None, None,
     '招き猫発祥の地のひとつとされる「今戸神社」は、縁結びのご利益で知られる神社です！境内には招き猫がたくさん並び、かわいらしい写真が撮れるスポットです。\n\n是非一度、招き猫に会いに「今戸神社」に訪れてみるのはいかがでしょうか？',
     'Imado Shrine, said to be one of the birthplaces of the maneki-neko (beckoning cat), is known for blessings in matchmaking! Many beckoning cats line the grounds, making it a charming photo spot.\n\nWhy not visit "Imado Shrine" to meet the beckoning cats?'),
    ('Ōtori Shrine', None, '鷲神社', 'Otori Shrine', ('観光', 'Tourism'), '史跡・寺社',
     ['# ひとり旅', '# 思い出', '# 穴場', '# 新発見', '# 大人っぽい'], None, None,
     '「鷲神社」は、11月の酉の日に開かれる「酉の市」で知られる神社です！縁起物の熊手を求める人々で賑わい、商売繁盛を願う参拝者が多く訪れます。\n\n是非一度、浅草の奥にある「鷲神社」に訪れてみるのはいかがでしょうか？',
     'Otori Shrine is famous for the Tori-no-ichi fair held on the days of the rooster in November! Visitors come to buy lucky kumade rakes and pray for business prosperity.\n\nWhy not visit "Otori Shrine", a short walk beyond central Asakusa?'),
    ('Shitaya Shrine', None, '下谷神社', 'Shitaya Shrine', ('観光', 'Tourism'), '史跡・寺社',
     ['# ひとり旅', '# 思い出', '# 穴場', '# 新発見', '# 大人っぽい'], None, None,
     '「下谷神社」は台東区で最も古い神社とされ、境内は江戸で最初の寄席が開かれた「寄席発祥の地」としても知られています！静かな境内で歴史を感じられるスポットです。\n\n是非一度、落語ゆかりの「下谷神社」に訪れてみるのはいかがでしょうか？',
     "Shitaya Shrine is said to be the oldest shrine in Taito City, and its grounds are known as the birthplace of yose, Edo's first rakugo halls! A quiet place to feel the area's history.\n\nWhy not visit \"Shitaya Shrine\", a shrine linked to rakugo?"),
    ('Torikoe Shrine', None, '鳥越神社', 'Torikoe Shrine', ('観光', 'Tourism'), '史跡・寺社',
     ['# ひとり旅', '# 思い出', '# 穴場', '# 新発見'], None, None,
     '「鳥越神社」は、6月の鳥越祭で担がれる東京有数の重さを誇る「千貫神輿」で知られる、歴史ある神社です！普段は地域に親しまれる落ち着いた雰囲気のスポットです。\n\n是非一度、下町の歴史を感じる「鳥越神社」に訪れてみるのはいかがでしょうか？',
     'Torikoe Shrine is a historic shrine known for its "Sengan Mikoshi", one of the heaviest portable shrines in Tokyo, carried at the June festival! On other days it is a calm, local spot.\n\nWhy not visit "Torikoe Shrine" to feel the history of shitamachi Tokyo?'),
    ('Kuramae Shrine', None, '蔵前神社', 'Kuramae Shrine', ('観光', 'Tourism'), '史跡・寺社',
     ['# ひとり旅', '# 思い出', '# 穴場', '# 新発見'], None, None,
     '「蔵前神社」は、江戸時代に勧進相撲が行われたことでも知られる八幡様をお祀りする神社です！相撲ゆかりの歴史を感じられる、落ち着いたスポットです。\n\n是非一度、相撲の歴史が残る「蔵前神社」に訪れてみるのはいかがでしょうか？',
     'Kuramae Shrine, dedicated to Hachiman, is known as a place where kanjin sumo tournaments were held in the Edo period! A quiet spot with a link to sumo history.\n\nWhy not visit "Kuramae Shrine" and its sumo history?'),
    ('Iriya South Park', None, '入谷南公園', 'Iriya South Park', ('観光', 'Tourism'), '公園',
     ['# 子供連れ', '# ひとり旅', '# 癒やし', '# 穴場'], ('24h', None), None,
     'かっぱ橋道具街の北、入谷の街なかにある「入谷南公園」は、地域の人々に親しまれる公園です！散策の途中にひと休みできるスポットです。\n\n是非一度、下町の日常を感じに「入谷南公園」に訪れてみるのはいかがでしょうか？',
     'Iriya South Park, north of Kappabashi in the Iriya neighbourhood, is a park loved by local residents! A good place for a break while exploring.\n\nWhy not visit "Iriya South Park" to see everyday life in shitamachi Tokyo?'),
    ('Yasaki Inari Shrine', None, '矢先稲荷神社', 'Yasaki Inari Shrine', ('観光', 'Tourism'), '史跡・寺社',
     ['# ひとり旅', '# 思い出', '# 穴場', '# 新発見'], None, None,
     'かっぱ橋道具街に建つ「矢先稲荷神社」は、拝殿の天井に日本の馬乗りの歴史を描いた天井画があることで知られる神社です！道具街散策のついでに立ち寄れるスポットです。\n\n是非一度、天井画が見事な「矢先稲荷神社」に訪れてみるのはいかがでしょうか？',
     'Yasaki Inari Shrine, on Kappabashi street, is known for the ceiling of its hall, painted with scenes from the history of horse riding in Japan! An easy stop while exploring Kappabashi.\n\nWhy not visit "Yasaki Inari Shrine" to see its ceiling paintings?'),
    ('Dairokuten Sakaki Shrine', None, '第六天榊神社', 'Dairokuten Sakaki Shrine', ('観光', 'Tourism'), '史跡・寺社',
     ['# ひとり旅', '# 穴場', '# 新発見', '# 大人っぽい'], None, None,
     '蔵前・柳橋エリアにある「第六天榊神社」は、古い歴史をもつ神社です！ビルの間に佇む静かな境内で、下町の歴史を感じられるスポットです。\n\n是非一度、街なかの静かな「第六天榊神社」に訪れてみるのはいかがでしょうか？',
     'Dairokuten Sakaki Shrine in the Kuramae / Yanagibashi area is a shrine with a long history! Its quiet grounds among the buildings let you feel the history of shitamachi Tokyo.\n\nWhy not visit the quiet "Dairokuten Sakaki Shrine"?'),
    ('World Bags & Luggage Museum', None, '世界のカバン博物館', 'World Bags & Luggage Museum', ('観光', 'Tourism'), '博物館・美術館',
     ['# 子供連れ', '# ひとり旅', '# 雨の日', '# 穴場', '# 新発見'], None, None,
     'カバンメーカーが運営する「世界のカバン博物館」では、世界各国のカバンや歴史的なトランクなどを無料で見学できます！雨の日にもおすすめのスポットです。\n\n是非一度、カバンの歴史をたどれる「世界のカバン博物館」に訪れてみるのはいかがでしょうか？',
     'At the World Bags & Luggage Museum, run by a luggage maker, you can see bags and historic trunks from around the world for free! A good choice on a rainy day.\n\nWhy not visit the "World Bags & Luggage Museum" to explore the history of bags?'),
    ('Edo Taitō Traditional Crafts Museum', None, '江戸たいとう伝統工芸館', 'Edo Taito Traditional Crafts Museum', ('観光', 'Tourism'), '博物館・美術館',
     ['# 子供連れ', '# ひとり旅', '# 雨の日', '# 穴場', '# 新発見'], None, None,
     '「江戸たいとう伝統工芸館」では、台東区で受け継がれてきた伝統工芸品を無料で見学でき、職人による実演が行われることもあります！ものづくりの技にふれられるスポットです。\n\n是非一度、江戸の技を感じに「江戸たいとう伝統工芸館」に訪れてみるのはいかがでしょうか？',
     'At the Edo Taito Traditional Crafts Museum you can see traditional crafts made in Taito City for free, sometimes with demonstrations by craftspeople! A place to see Edo craftsmanship up close.\n\nWhy not visit the "Edo Taito Traditional Crafts Museum"?'),
    ('Asakusa Fuji Sengen-jinja Shrine', None, '浅草富士浅間神社', 'Asakusa Fuji Sengen Shrine', ('観光', 'Tourism'), '史跡・寺社',
     ['# ひとり旅', '# 穴場', '# 新発見', '# 癒やし'], None, None,
     '富士山信仰の神社「浅草富士浅間神社」は、初夏に開かれる植木市でも知られています！浅草寺の北にある、地域に親しまれる小さな神社です。\n\n是非一度、浅草の北にある「浅草富士浅間神社」に訪れてみるのはいかがでしょうか？',
     'Asakusa Fuji Sengen Shrine, a shrine of Mt Fuji worship, is also known for its plant fair in early summer! A small local shrine north of Senso-ji.\n\nWhy not visit "Asakusa Fuji Sengen Shrine", just north of Senso-ji?'),
    ('Sumida River Ryokudō Park', None, '隅田川緑道公園', 'Sumida River Ryokudo Park', ('観光', 'Tourism'), '公園',
     ['# ひとり旅', '# デート', '# 癒やし', '# 穴場'], ('24h', None), None,
     '隅田川沿いに続く「隅田川緑道公園」は、川風を感じながら散歩できる緑道です！蔵前の街なかから少し歩くだけで、川沿いの開放的な景色を楽しめるスポットです。\n\n是非一度、川沿いを散策しに「隅田川緑道公園」に訪れてみるのはいかがでしょうか？',
     'Sumida River Ryokudo Park is a green walkway along the Sumida River where you can stroll in the river breeze! Just a short walk from Kuramae, it offers open riverside views.\n\nWhy not take a riverside walk at "Sumida River Ryokudo Park"?'),
    ('Miyamoto-Unosuke Drum Museum', None, '太鼓館', 'Drum Museum (Taiko-kan)', ('観光', 'Tourism'), '博物館・美術館',
     ['# 子供連れ', '# ひとり旅', '# 雨の日', '# 穴場', '# 新発見'], None, None,
     '太鼓店が運営する「太鼓館」では、世界各国の太鼓を展示しており、実際に叩いて音を楽しめる太鼓もあります！子どもから大人まで楽しめる体験型のスポットです。\n\n是非一度、世界の太鼓にふれに「太鼓館」に訪れてみるのはいかがでしょうか？',
     'The Drum Museum (Taiko-kan), run by a drum maker, displays drums from around the world, and you can play some of them yourself! A hands-on spot for all ages.\n\nWhy not visit the "Drum Museum" to try drums from around the world?'),
]

NEW_ID = 1204  # next free ID after the current maximum (1203); confirm with Le Laboratory
wb = Workbook()
wj = wb.active; wj.title = 'spotsinfo'
we = wb.create_sheet('spotsinfo - EN')
FONT = Font(name='Arial', size=10)
for ws, hdr in ((wj, HDR_JP), (we, HDR_EN)):
    ws.append(hdr)
nid = NEW_ID
for key, sid, jp, en, (cj, ce), gj, tags, hrs, url, cmt_jp, cmt_en in SPOTS:
    lat, lng = coords[key]
    if sid is None:
        sid, nid = nid, nid + 1
    hj, he = hours(*(hrs or (None, None)))
    u = url or gmaps(lat, lng)
    wj.append([sid, jp, lat, lng, cj, gj, ','.join(tags), 'あり', LOC_JP, hj, u, cmt_jp])
    we.append([sid, en, lat, lng, ce, GENRE[gj], ', '.join(TAG[t] for t in tags), 'Yes', LOC_EN, he, u, cmt_en])
for ws in (wj, we):
    for row in ws.iter_rows():
        for c in row:
            c.font = FONT
            c.alignment = Alignment(vertical='top', wrap_text=c.column in (7, 10, 12))
    for col, w in zip('ABCDEFGHIJKL', [8, 32, 12, 12, 10, 24, 40, 9, 26, 40, 40, 80]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'C2'

# ---- Stamp rally sheet (fields used by the app's StampRally / Goal models) ----
sr = wb.create_sheet('Stamp rally')
sr.append(['Field (app model)', 'Proposed value', 'Notes'])
for r in [
    ['title', '浅草チェックポイントラリー / Asakusa Checkpoint Rally', ''],
    ['summary', '16か所のチェックポイントでスタンプを集めよう！ / Collect stamps at 16 checkpoints!', ''],
    ['start_at', '2026-11-28', 'Pilot day; set to cover all sessions'],
    ['end_at', '2027-01-30', 'Last (rain) date'],
    ['event_url', '', 'Optional page explaining the rally'],
    ['spots', 'All 16 spots in this file', 'Each spot gets event_stamp = this rally'],
]:
    sr.append(r)
sr.append([])
sr.append(['Goal index (stamps)', 'bonus_point', 'bonus_item'])
for g, p in [(4, 30), (8, 60), (12, 90), (16, 150)]:
    sr.append([g, p, ''])
sr.append([])
sr.append(['bonus_item is left empty on purpose: a physical prize would add an incentive, and claiming it asks for an e-mail address.'])
for row in sr.iter_rows():
    for c in row:
        c.font = FONT
for col, w in zip('ABC', [26, 70, 60]):
    sr.column_dimensions[col].width = w

# ---- Notes sheet ----
nt = wb.create_sheet('Notes')
notes = [
    ['Item', 'Note'],
    ['Format', "Same 12 columns as suisuitabi_information_-_EN.xlsx (sheets 'spotsinfo' and 'spotsinfo - EN'); the header 'Longtitude' is kept as in the original."],
    ['IDs', 'Senso-ji reuses existing ID 1041 (Mission changed to あり / Yes). The other 15 spots use new IDs 1204-1218, after the current maximum (1203). Confirm with Le Laboratory.'],
    ['Location', "上野・浅草・日暮里・谷根千, as for Senso-ji (1041) and Sumida Park (1056). Their English rows say 'Ageo, Kuki, Kōnosu, Ina', which looks like a translation error; here it is 'Ueno, Asakusa, Nippori, Yanesen'."],
    ['Mission', 'あり / Yes for all 16: one photo mission each (gamification group). The control-group app hides missions.'],
    ['Coordinates', 'From Google Maps (via Apify), collected 13 Oct 2026.'],
    ['Business hours', 'Filled only where known: Kaminarimon, parks (24 hours), Senso-ji main hall (6:30-17:00 in October-March; 6:00 from April). Left empty for shrines, museums and Kappabashi shops: confirm on site or official sites.'],
    ['URL', 'Official sites for Senso-ji and Kappabashi. The others link to the Google Maps location; replace with official pages where they exist.'],
    ['Recommendation points', 'Written in the same style as the existing entries. Senso-ji keeps its existing text. Facts are from general knowledge: check before publishing.'],
    ['Names to check', 'Sumida River Ryokudo Park (隅田川緑道公園) is the Google Maps name; confirm the official Japanese name.'],
    ['Tags', "Less popular spots carry '# 穴場' (Hidden spots). Both groups see the same tags, so this does not differ between groups."],
]
for r in notes:
    nt.append(r)
for row in nt.iter_rows():
    for c in row:
        c.font = FONT
        c.alignment = Alignment(vertical='top', wrap_text=True)
nt.column_dimensions['A'].width = 22
nt.column_dimensions['B'].width = 110
for ws in (wj, we, sr, nt):
    for c in ws[1]:
        c.font = Font(name='Arial', size=10, bold=True)
wb.save('asakusa_spots_suisuitabi_format.xlsx')
print('saved', len(SPOTS), 'spots; new IDs up to', nid - 1)
