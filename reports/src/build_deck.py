import json
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x0F, 0x2F, 0x46)
TEAL = RGBColor(0x7B, 0xCF, 0xDF)
TINT = RGBColor(0xE9, 0xF6, 0xF9)
TEXT = RGBColor(0x26, 0x2B, 0x30)
MUTED = RGBColor(0x5A, 0x64, 0x6E)
ORANGE = RGBColor(0xC2, 0x5E, 0x12)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Arial'

# content area (EMU) inside the template's header line and footer bar
L, R = 900_000, 12_550_000
T, B = 1_380_000, 6_900_000
W = R - L

prs = Presentation('base.pptx')
slides = list(prs.slides)


def runs(p, parts, size, color=TEXT, bold=False):
    # parts: str or list of (text, bold) tuples
    if isinstance(parts, str):
        parts = [(parts, bold)]
    for text, b in parts:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name, f.size, f.bold, f.color.rgb = FONT, Pt(size), b, color


def text(slide, x, y, w, h, paras, size=16, color=TEXT, anchor=MSO_ANCHOR.TOP,
         align=PP_ALIGN.LEFT, space=6, name=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, 0)
    first = True
    for para in paras:
        if isinstance(para, dict):
            spec = para
        else:
            spec = {'t': para}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = spec.get('align', align)
        p.space_after = Pt(spec.get('space', space))
        if spec.get('bullet'):
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', '228600')
            pPr.set('indent', '-228600')
            bu = pPr.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar', {'char': '•'})
            pPr.append(bu)
        runs(p, spec['t'], spec.get('size', size), spec.get('color', color), spec.get('bold', False))
    return tb


def box(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ''
    return s


def set_title(slide, title):
    for sh in list(slide.shapes):
        if sh.name == 'Google Shape;52;p8':
            sh.width = Emu(9_800_000)
            sh.text_frame.paragraphs[0].runs[0].text = title
        elif sh.name == 'Google Shape;53;p8':
            sh._element.getparent().remove(sh._element)


def table(slide, x, y, w, rows, widths, size=14, header_fill=NAVY, row_h=380_000, highlight=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Emu(x), Emu(y), Emu(w), Emu(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    # drop the default table style banding so our fills apply
    tblPr.set('bandRow', '0')
    tblPr.set('firstRow', '1')
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Emu(int(w * cw))
    for r, row in enumerate(rows):
        tbl.rows[r].height = Emu(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Emu(70_000)
            cell.margin_top = cell.margin_bottom = Emu(30_000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ''
            is_num = isinstance(val, (int, float))
            p.alignment = PP_ALIGN.RIGHT if is_num else PP_ALIGN.LEFT
            sval = f'{val:,}' if is_num else str(val)
            if r == 0:
                runs(p, sval, size, WHITE, True)
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            else:
                hl = highlight and highlight(row)
                runs(p, sval, size, ORANGE if hl else TEXT, bool(hl) and c == 1)
                cell.fill.solid(); cell.fill.fore_color.rgb = TINT if r % 2 == 0 else WHITE
    return shp


def arrow(slide, x, y, w, h):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(x), Emu(y), Emu(w), Emu(h))
    a.fill.solid(); a.fill.fore_color.rgb = TEAL
    a.line.fill.background(); a.shadow.inherit = False
    return a


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ---------- 1. Cover ----------
s = slides[0]
text(s, L, 1_250_000, 11_200_000, 1_500_000,
     [{'t': 'Spot Selection for a Gamified Tourism Experiment in Asakusa', 'size': 36, 'bold': True, 'color': NAVY}],
     name='Cover title')
text(s, L, 2_900_000, 11_200_000, 600_000,
     [{'t': 'Progress report: study area, checkpoints and open decisions', 'size': 20, 'color': NAVY}],
     name='Cover subtitle')
text(s, L, 4_950_000, 11_200_000, 900_000,
     [{'t': 'HyoSeok (Jake) Choi', 'size': 20, 'bold': True, 'color': NAVY},
      {'t': '13 October 2026', 'size': 16, 'color': MUTED}], name='Cover byline')
notes(s, 'Progress report on choosing the study area and the experiment checkpoints, and the decisions still open.')

# ---------- 2. Table of contents ----------
s = slides[1]
set_title(s, 'Table of contents')
toc = ['Background and research goal', 'Study area: why Asakusa', 'Spot selection method',
       'Selected checkpoints', 'Open decisions on the spot list', 'Experiment design',
       'Hypotheses and measures', 'Questionnaires and app changes', 'Next steps and questions']
for i, t_ in enumerate(toc):
    y = T + 60_000 + i * 600_000
    c = box(s, L, y, 520_000, 520_000, fill=TEAL if i % 2 == 0 else NAVY, shape=MSO_SHAPE.OVAL, name=f'TOC number {i+1}')
    text(s, L, y, 520_000, 520_000, [{'t': str(i + 1), 'size': 18, 'bold': True,
         'color': NAVY if i % 2 == 0 else WHITE, 'align': PP_ALIGN.CENTER, 'space': 0}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, L + 750_000, y, 9_000_000, 520_000, [{'t': t_, 'size': 20, 'space': 0}], anchor=MSO_ANCHOR.MIDDLE)

# ---------- 3. Background ----------
s = slides[2]
set_title(s, 'Background and research goal')
box(s, L, T, 5_300_000, 2_350_000, fill=NAVY, name='Research question card')
text(s, L + 280_000, T + 230_000, 4_750_000, 1_950_000,
     [{'t': 'Research question', 'size': 16, 'bold': True, 'color': TEAL},
      {'t': 'Can gamification features in a tourism app lead visitors from crowded famous spots to less popular spots nearby?', 'size': 20, 'color': WHITE}])
text(s, L, T + 2_600_000, 5_300_000, 2_800_000,
     [{'t': 'Reference study', 'size': 16, 'bold': True, 'color': NAVY},
      {'t': 'Kawanaka et al. (2020), Smart Cities 3, 736-757', 'size': 14, 'color': MUTED},
      {'t': 'Kyoto, 33 students, 3 reward schemes', 'size': 16, 'bullet': True},
      {'t': 'Higher points at lesser-known spots changed routes and stay time', 'size': 16, 'bullet': True},
      {'t': 'Check-in missions lowered sightseeing enjoyment', 'size': 16, 'bullet': True}])
rows = [['', 'Kawanaka et al. (2020)', 'This study'],
        ['Framing', 'Collecting tourism data', 'Reducing overtourism'],
        ['Control group', 'None', 'Yes: with vs without gamification'],
        ['Participants', '33 students, mostly Japanese', 'Japanese and foreign tourists'],
        ['Comparison', 'User motivation type', 'Nationality and region'],
        ['Area', 'Kyoto (Higashiyama)', 'Tokyo (Asakusa)']]
text(s, 6_550_000, T, 6_000_000, 400_000, [{'t': 'What this study adds', 'size': 16, 'bold': True, 'color': NAVY}])
table(s, 6_550_000, T + 480_000, 6_000_000, rows, [0.27, 0.36, 0.37], size=14, row_h=560_000)
notes(s, 'The closest prior work is Kawanaka et al. 2020. Our study differs mainly in framing, a control group, and a mixed Japanese and foreign sample.')

# ---------- 4. Study area ----------
s = slides[3]
set_title(s, 'Study area: why Asakusa')
cards = [('Clear overtourism hotspot', 'Sensō-ji has about 100,000 Google reviews, the most in the area by far.'),
         ('Japanese and foreign visitors', 'Needed for the nationality comparison in this study.'),
         ('Alternatives nearby', 'Shrines, parks, museums and shopping streets 0.5–1.5 km from the core.'),
         ('Walkable', 'Everything fits in a 1.5 km radius, about a 20-minute walk from Kaminarimon.')]
cw = (W - 3 * 300_000) // 4
for i, (h_, b_) in enumerate(cards):
    x = L + i * (cw + 300_000)
    box(s, x, T, cw, 2_700_000, name=f'Criterion card {i+1}')
    text(s, x + 200_000, T + 200_000, 500_000, 500_000, [{'t': f'0{i+1}', 'size': 28, 'bold': True, 'color': TEAL}])
    text(s, x + 200_000, T + 800_000, cw - 400_000, 1_800_000,
         [{'t': h_, 'size': 20, 'bold': True, 'color': NAVY}, {'t': b_, 'size': 17}])
box(s, L, T + 3_050_000, W, 2_150_000, fill=WHITE, line=TEAL, name='Boundary card')
text(s, L + 300_000, T + 3_250_000, W - 600_000, 1_800_000,
     [{'t': 'Study area boundary', 'size': 18, 'bold': True, 'color': NAVY},
      {'t': 'Circle of 1.5 km radius centred on Kaminarimon Gate', 'size': 16, 'bullet': True},
      {'t': 'West bank of the Sumida River only, so Tokyo Skytree does not act as a second hub', 'size': 16, 'bullet': True},
      {'t': 'Less popular spots must be 300 m to 1.5 km from a popular spot: close enough to substitute, far enough to be a separate place', 'size': 16, 'bullet': True}])

# ---------- 5. Method ----------
s = slides[4]
set_title(s, 'Spot selection method')
steps = [('520', 'places', 'Google Maps scrape within 1.5 km of Kaminarimon (13 Oct 2026)'),
         ('56', 'candidates', 'Public, non-commercial, west of the river, 50+ reviews'),
         ('3', 'classes', 'Jenks natural breaks on log review count'),
         ('16', 'checkpoints', 'All at least 300 m apart (app check-in radius)')]
sw, gap = 2_450_000, (W - 4 * 2_450_000) // 3
for i, (n, u, d) in enumerate(steps):
    x = L + i * (sw + gap)
    box(s, x, T + 200_000, sw, 3_300_000, fill=NAVY if i == 3 else TINT, name=f'Step {i+1}')
    col = WHITE if i == 3 else NAVY
    text(s, x + 200_000, T + 400_000, sw - 400_000, 2_900_000,
         [{'t': f'Step {i+1}', 'size': 14, 'bold': True, 'color': TEAL if i == 3 else MUTED},
          {'t': n, 'size': 54, 'bold': True, 'color': col, 'space': 0},
          {'t': u, 'size': 18, 'bold': True, 'color': col, 'space': 14},
          {'t': d, 'size': 15, 'color': WHITE if i == 3 else TEXT}])
    if i < 3:
        arrow(s, x + sw + (gap - 320_000) // 2, T + 1_650_000, 320_000, 380_000)
text(s, L, T + 3_850_000, W, 1_500_000,
     [{'t': 'Why review counts?', 'size': 16, 'bold': True, 'color': NAVY},
      {'t': 'Reproducible and widely used as a proxy for visitor volume. Kawanaka et al. (2020) used a similar measure (web search hits).', 'size': 15, 'bullet': True},
      {'t': 'Jenks breaks set the popular / less popular cut at a natural gap in the data, not an arbitrary percentile.', 'size': 15, 'bullet': True}])

# ---------- 6. Map ----------
s = slides[5]
set_title(s, 'Result: 16 checkpoints')
mh = B - T
s.shapes.add_picture('/home/user/GitHub/spot-selection/asakusa_spots_map.png', Emu(L), Emu(T - 50_000), Emu(mh), Emu(mh))
rx = L + mh + 350_000
rw = R - rx
cls = [('High (popular)', '≥ 13,091 reviews', 'Sensō-ji, Kaminarimon, Kappabashi', NAVY, WHITE),
       ('Medium', '597 – 13,090 reviews', '7 checkpoints, mostly shrines', TINT, NAVY),
       ('Low', '< 597 reviews', '6 checkpoints: museums, parks, shrines', TINT, NAVY)]
for i, (a, b_, c_, f, col) in enumerate(cls):
    y = T + i * 1_300_000
    box(s, rx, y, rw, 1_150_000, fill=f, name=f'Class {a}')
    text(s, rx + 220_000, y + 150_000, rw - 440_000, 900_000,
         [{'t': [(a + '  ', True), (b_, False)], 'size': 16, 'color': col},
          {'t': c_, 'size': 14, 'color': col if f == NAVY else TEXT}])
text(s, rx, T + 3_950_000, rw, 1_500_000,
     [{'t': 'Nakamise is merged into the Kaminarimon checkpoint (85 m apart).', 'size': 14, 'bullet': True},
      {'t': 'Minimum distance between any two checkpoints: 310 m.', 'size': 14, 'bullet': True},
      {'t': 'Distances are straight-line, not walking routes.', 'size': 14, 'bullet': True, 'color': MUTED}])

# ---------- 7. Checkpoint table ----------
s = slides[6]
set_title(s, 'Selected checkpoints')
d = json.load(open('/home/user/GitHub/spot-selection/final.json'))
sel = [r for r in d['rows'] if r['selected']]
hdr = ['#', 'Spot', 'Type', 'Reviews', 'Dist. (m)']
body = [[i + 1, r['name'].replace(' Kitchen Supply Town', ' Kitchen Town').replace(' Traditional Crafts Museum', ' Crafts Museum').replace(' Shrine', '').replace('-jinja', ''),
         r['type'].replace(' (free)', ''), r['reviews'], '–' if r['label'] == 'Popular' else r['dist_nearest_popular_m']]
        for i, r in enumerate(sel)]
half = (W - 350_000) // 2
is_pop = lambda row: row[4] == '–'
table(s, L, T, half, [hdr] + body[:8], [0.07, 0.42, 0.19, 0.16, 0.16], size=13, row_h=520_000, highlight=is_pop)
table(s, L + half + 350_000, T, half, [hdr] + body[8:], [0.07, 0.42, 0.19, 0.16, 0.16], size=13, row_h=520_000)
text(s, L, T + 5_320_000, W, 300_000,
     [{'t': 'Orange = popular checkpoint. Dist. = straight-line distance to the nearest popular checkpoint. "Shrine" dropped from names to save space; reviews = Google Maps review count on 13 Oct 2026.', 'size': 12, 'color': MUTED}])

# ---------- 8. Open decisions ----------
s = slides[7]
set_title(s, 'Open decisions on the spot list')
dec = [('1. Spot variety', '9 of 13 less popular spots are shrines.',
        'Swap Yasaki Inari → Kappabashi Hondōri, and Kuramae Shrine → Seika Park (7 of 13).'),
       ('2. Shōden or Imado Shrine', 'Only 209 m apart, so one must go.',
        'Keep Imado (4,271 vs 2,495 reviews), following the stated rule.'),
       ('3. Radius: 1.5 or 1.2 km', '1.2 km removes 4 outer shrines, ~21 min walk each way.',
        '1.2 km for 2–3 h sessions; 1.5 km for 4 h or more.'),
       ('4. River boundary', 'The river line was traced by hand.',
        'Check addresses of 3 spots near the river (should say 台東区).')]
cw2, ch2 = (W - 300_000) // 2, 2_550_000
for i, (h_, issue, rec) in enumerate(dec):
    x = L + (i % 2) * (cw2 + 300_000)
    y = T + (i // 2) * (ch2 + 250_000)
    box(s, x, y, cw2, ch2, name=f'Decision card {i+1}')
    text(s, x + 250_000, y + 200_000, cw2 - 500_000, ch2 - 300_000,
         [{'t': h_, 'size': 22, 'bold': True, 'color': NAVY, 'space': 10},
          {'t': issue, 'size': 18, 'space': 14},
          {'t': [('Recommendation: ', True), (rec, False)], 'size': 18, 'color': NAVY}])

# ---------- 9. Experiment design ----------
s = slides[8]
set_title(s, 'Experiment design: with vs without gamification')
rows = [['Group', 'What participants see in the app'],
        ['Without gamification (control)', 'Map with all 16 checkpoints and spot information'],
        ['With gamification', 'The same map, plus points, missions, a stamp rally, ranking and medals']]
table(s, L, T, 7_300_000, rows, [0.34, 0.66], size=15, row_h=700_000)
text(s, L, T + 2_350_000, 7_300_000, 3_200_000,
     [{'t': 'Only gamification differs', 'size': 20, 'bold': True, 'color': NAVY},
      {'t': [('Equal rewards at every checkpoint: ', True), ('no extra points for less popular spots', False)], 'size': 17, 'bullet': True},
      {'t': [('Stamp rally of all 16 checkpoints: ', True), ('13 of 16 stamps are at less popular spots', False)], 'size': 17, 'bullet': True},
      {'t': [('Same fee for everyone: ', True), ('no bonus by ranking', False)], 'size': 17, 'bullet': True},
      {'t': [('Same measurement: ', True), ('automatic GPS check-in in both groups', False)], 'size': 17, 'bullet': True},
      {'t': [('Random assignment ', True), ('within Japanese and foreign participants', False)], 'size': 17, 'bullet': True}])
sx, swd = L + 7_650_000, R - (L + 7_650_000)
stats = [('20', 'per group', 'Planned: 40 total; detects d ≈ 0.9 (large)'),
         ('15', 'per group', 'Minimum: 30 total; detects d ≈ 1.1 (very large)'),
         ('11', 'per group', 'Kawanaka et al. (2020), for comparison')]
for i, (n, u, d_) in enumerate(stats):
    y = T + i * 1_800_000
    box(s, sx, y, swd, 1_600_000, fill=NAVY if i < 2 else TINT, name=f'Sample size {n}')
    col = WHITE if i < 2 else NAVY
    text(s, sx + 250_000, y + 150_000, swd - 500_000, 1_350_000,
         [{'t': [(n + ' ', True), (u, False)], 'size': 28, 'color': col, 'space': 4},
          {'t': d_, 'size': 16, 'color': col if i < 2 else TEXT}])
text(s, sx, T + 5_420_000, swd, 300_000,
     [{'t': 'α = 0.05, power = 0.80; to be confirmed with G*Power. Small sample, so results are exploratory.', 'size': 12, 'color': MUTED}])

# ---------- 9b. Hypotheses ----------
s = slides[9]
set_title(s, 'Hypotheses and measures')
rows = [['#', 'The gamification group…', 'Measure (app logs)', 'Test'],
        ['H1', 'visits more less popular checkpoints', 'Distinct less popular checkpoints visited', 'Mann–Whitney U'],
        ['H2', 'has a larger share of visits at less popular checkpoints', 'Less popular visits ÷ all checkpoint visits', 'Mann–Whitney U'],
        ['H3', 'travels farther from the crowded core', 'Max distance from Kaminarimon; time > 300 m from popular spots', 'Mann–Whitney U'],
        ['H4', 'shows an effect that differs by origin', 'H1 measure by group × Japanese / foreign', 'Exploratory (descriptive only)'],
        ['H5', 'is no less satisfied with the trip', 'Post-questionnaire satisfaction', 'Mann–Whitney U']]
table(s, L, T, W, rows, [0.06, 0.32, 0.38, 0.24], size=15, row_h=640_000)
box(s, L, T + 4_250_000, W, 1_150_000, name='H1 H2 note')
text(s, L + 280_000, T + 4_250_000, W - 560_000, 1_150_000,
     [{'t': [('Why H1 and H2 together: ', True), ('gamification may raise visits everywhere. H1 alone could rise with no shift away from popular spots; H2 shows whether the balance actually moves.', False)], 'size': 16, 'color': NAVY, 'space': 0}],
     anchor=MSO_ANCHOR.MIDDLE)

# ---------- 10. Surveys and app changes ----------
s = slides[10]
set_title(s, 'Questionnaires and app changes')
colw = (W - 350_000) // 2
box(s, L, T, colw, 5_450_000, name='Survey card')
text(s, L + 280_000, T + 230_000, colw - 560_000, 5_000_000,
     [{'t': 'Questionnaires (separate from the app)', 'size': 20, 'bold': True, 'color': NAVY, 'space': 4},
      {'t': 'Paper or online form, linked to app logs by participant ID', 'size': 15, 'color': MUTED, 'space': 14},
      {'t': 'Pre (before the walk)', 'size': 18, 'bold': True, 'color': NAVY},
      {'t': 'Age, gender, nationality, country of residence', 'size': 16, 'bullet': True},
      {'t': 'Previous visits; checkpoints already known', 'size': 16, 'bullet': True},
      {'t': 'Travel motivation; experience with games and apps', 'size': 16, 'bullet': True, 'space': 14},
      {'t': 'Post (right after the walk)', 'size': 18, 'bold': True, 'color': NAVY},
      {'t': 'Satisfaction, discovery of new places, perceived crowding, revisit intention', 'size': 16, 'bullet': True},
      {'t': 'Why they chose the places they visited (free answer)', 'size': 16, 'bullet': True},
      {'t': 'Gamification group: which game features affected their route', 'size': 16, 'bullet': True}])
box(s, L + colw + 350_000, T, colw, 5_450_000, fill=WHITE, line=TEAL, name='App card')
text(s, L + colw + 630_000, T + 230_000, colw - 560_000, 5_000_000,
     [{'t': 'App changes (from the code review)', 'size': 20, 'bold': True, 'color': NAVY},
      {'t': [('Control-group switch: ', True), ('a flag exists but is unused; hide points, missions, ranking and medals', False)], 'size': 17, 'bullet': True},
      {'t': [('Stamp rally and missions ', True), ('at all 16 checkpoints; equal points everywhere (current 10 per check-in is fine)', False)], 'size': 17, 'bullet': True},
      {'t': [('Popular / less popular tag ', True), ('on each spot, to count visits by type', False)], 'size': 17, 'bullet': True},
      {'t': [('Add the 16 checkpoints ', True), ('to the server database', False)], 'size': 17, 'bullet': True},
      {'t': [('Participant ID ', True), ('at registration; no in-app survey changes needed', False)], 'size': 17, 'bullet': True}])

# ---------- 11. Next steps ----------
s = slides[11]
set_title(s, 'Next steps and questions')
box(s, L, T, colw, 5_450_000, name='Next steps card')
text(s, L + 280_000, T + 230_000, colw - 560_000, 5_000_000,
     [{'t': 'Next steps', 'size': 20, 'bold': True, 'color': NAVY},
      {'t': 'Finalise the spot list (decisions 1–4)', 'size': 17, 'bullet': True},
      {'t': 'Count pedestrians at each checkpoint on one weekend to confirm crowding', 'size': 17, 'bullet': True},
      {'t': 'Agree app changes with Le Laboratory', 'size': 17, 'bullet': True},
      {'t': 'Prepare the pre- and post-questionnaires', 'size': 17, 'bullet': True, 'space': 16},
      {'t': 'Timing', 'size': 20, 'bold': True, 'color': NAVY},
      {'t': 'Weekend afternoons; all groups on the same days', 'size': 17, 'bullet': True},
      {'t': 'Avoid Sanja Matsuri (May), Sumida fireworks (July), Tori-no-ichi (Nov), New Year', 'size': 17, 'bullet': True},
      {'t': 'Plan a rain date', 'size': 17, 'bullet': True}])
box(s, L + colw + 350_000, T, colw, 5_450_000, fill=NAVY, name='Questions card')
text(s, L + colw + 630_000, T + 230_000, colw - 560_000, 5_000_000,
     [{'t': 'Questions for you', 'size': 20, 'bold': True, 'color': TEAL},
      {'t': '1. Is 1.2 km or 1.5 km more suitable for our session length?', 'size': 19, 'color': WHITE, 'space': 12},
      {'t': '2. Is 15–20 per group (30–40 total) acceptable as an exploratory study?', 'size': 19, 'color': WHITE, 'space': 12},
      {'t': '3. When should the experiment run?', 'size': 19, 'color': WHITE, 'space': 12},
      {'t': '4. Is excluding commercial spots the right rule, or should a few free-entry shops be allowed?', 'size': 19, 'color': WHITE}])

prs.save('Asakusa_progress_report.pptx')
print('saved')
