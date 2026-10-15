import json
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x0F, 0x2F, 0x46)
TEAL = RGBColor(0x7B, 0xCF, 0xDF)
TINT = RGBColor(0xE9, 0xF6, 0xF9)
TEXT = RGBColor(0x26, 0x2B, 0x30)
MUTED = RGBColor(0x5A, 0x64, 0x6E)
ORANGE = RGBColor(0xC2, 0x5E, 0x12)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Arial'

# content area (EMU) inside the template's header line and footer bar
L, R = 900_000, 12_550_000
T, B = 1_380_000, 6_900_000
W = R - L

prs = Presentation(BASE)
slides = list(prs.slides)


def runs(p, parts, size, color=TEXT, bold=False):
    # parts: str or list of (text, bold) tuples
    if isinstance(parts, str):
        parts = [(parts, bold)]
    for text, b in parts:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name, f.size, f.bold, f.color.rgb = FONT, Pt(size), b, color


def text(slide, x, y, w, h, paras, size=16, color=TEXT, anchor=MSO_ANCHOR.TOP,
         align=PP_ALIGN.LEFT, space=6, name=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, 0)
    first = True
    for para in paras:
        if isinstance(para, dict):
            spec = para
        else:
            spec = {'t': para}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = spec.get('align', align)
        p.space_after = Pt(spec.get('space', space))
        if spec.get('bullet'):
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', '228600')
            pPr.set('indent', '-228600')
            bu = pPr.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar', {'char': '•'})
            pPr.append(bu)
        runs(p, spec['t'], spec.get('size', size), spec.get('color', color), spec.get('bold', False))
    return tb


def box(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ''
    return s


def set_title(slide, title):
    for sh in list(slide.shapes):
        if sh.name == 'Google Shape;52;p8':
            sh.width = Emu(9_800_000)
            sh.text_frame.paragraphs[0].runs[0].text = title
        elif sh.name == 'Google Shape;53;p8':
            sh._element.getparent().remove(sh._element)


def table(slide, x, y, w, rows, widths, size=14, header_fill=NAVY, row_h=380_000, highlight=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Emu(x), Emu(y), Emu(w), Emu(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    # drop the default table style banding so our fills apply
    tblPr.set('bandRow', '0')
    tblPr.set('firstRow', '1')
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Emu(int(w * cw))
    for r, row in enumerate(rows):
        tbl.rows[r].height = Emu(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Emu(70_000)
            cell.margin_top = cell.margin_bottom = Emu(30_000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ''
            is_num = isinstance(val, (int, float))
            p.alignment = PP_ALIGN.RIGHT if is_num else PP_ALIGN.LEFT
            sval = f'{val:,}' if is_num else str(val)
            if r == 0:
                runs(p, sval, size, WHITE, True)
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            else:
                hl = highlight and highlight(row)
                runs(p, sval, size, ORANGE if hl else TEXT, bool(hl) and c == 1)
                cell.fill.solid(); cell.fill.fore_color.rgb = TINT if r % 2 == 0 else WHITE
    return shp


def arrow(slide, x, y, w, h):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(x), Emu(y), Emu(w), Emu(h))
    a.fill.solid(); a.fill.fore_color.rgb = TEAL
    a.line.fill.background(); a.shadow.inherit = False
    return a


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


